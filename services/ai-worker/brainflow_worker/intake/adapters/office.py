"""DOCX / PPTX / XLSX adapters (optional libraries). Macros never executed."""

from __future__ import annotations

import io
import zipfile
from typing import Any

from brainflow_worker.intake.adapters.context import AdapterContext
from brainflow_worker.intake.bundle import empty_bundle, issue, truncate_text, apply_security
from brainflow_worker.intake.quarantine import quarantine_path_name, merge_quarantine


DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_MIME = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _scan_ooxml_macros(data: bytes, ctx: AdapterContext) -> None:
    """Flag and quarantine VBA parts inside OOXML without executing them."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            for name in zf.namelist():
                q = quarantine_path_name(name)
                if q.macros_blocked or q.blocked:
                    ctx.quarantine = merge_quarantine(ctx.quarantine, q)
    except zipfile.BadZipFile:
        pass


class DocxAdapter:
    id = "docx"

    def supports(self, mime: str, filename: str | None, data: bytes) -> bool:
        if mime == DOCX_MIME:
            return True
        lower = (filename or "").lower()
        return lower.endswith(".docx")

    def extract(self, data: bytes, ctx: AdapterContext) -> dict[str, Any]:
        bundle = empty_bundle(
            file_id=ctx.file_id,
            content_hash=ctx.content_hash,
            mime_type=DOCX_MIME,
            size_bytes=ctx.size_bytes,
            path=ctx.path,
            path_mode=ctx.path_mode,
            adapter_id=self.id,
            sniffed_mime=ctx.sniffed_mime,
        )
        _scan_ooxml_macros(data, ctx)
        if ctx.quarantine.macros_blocked:
            bundle["warnings"].append(
                issue("Quarantined", "Macro-enabled payload detected; macros not executed")
            )

        try:
            from docx import Document  # type: ignore
        except ImportError:
            bundle["errors"].append(
                issue("NeedsAdapter", "python-docx not installed; install intake extras")
            )
            return _done(bundle, ctx, 0.0)

        try:
            doc = Document(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001
            bundle["errors"].append(issue("Corrupt", f"DOCX open failed: {exc}"))
            return _done(bundle, ctx, 0.0)

        order = 0
        for i, para in enumerate(doc.paragraphs):
            text = (para.text or "").strip()
            if not text:
                continue
            text, truncated = truncate_text(text, ctx.limits.max_segment_chars)
            if truncated:
                bundle["warnings"].append(issue("Truncated", f"paragraph {i} truncated"))
            bundle["segments"].append(
                {
                    "id": f"p-{i}",
                    "kind": "text",
                    "text": text,
                    "order": order,
                    "locator": {"kind": "docx_para", "paragraph_index": i},
                    "confidence": 0.95,
                }
            )
            order += 1
            if order >= ctx.limits.max_segments:
                bundle["warnings"].append(issue("ResourceLimit", "DOCX segment limit"))
                break
        if doc.core_properties and doc.core_properties.title:
            bundle["metadata"]["title"] = str(doc.core_properties.title)
        bundle["confidence"] = 0.95 if bundle["segments"] else 0.4
        return _done(bundle, ctx, bundle["confidence"])


class PptxAdapter:
    id = "pptx"

    def supports(self, mime: str, filename: str | None, data: bytes) -> bool:
        if mime == PPTX_MIME:
            return True
        return (filename or "").lower().endswith(".pptx")

    def extract(self, data: bytes, ctx: AdapterContext) -> dict[str, Any]:
        bundle = empty_bundle(
            file_id=ctx.file_id,
            content_hash=ctx.content_hash,
            mime_type=PPTX_MIME,
            size_bytes=ctx.size_bytes,
            path=ctx.path,
            path_mode=ctx.path_mode,
            adapter_id=self.id,
            sniffed_mime=ctx.sniffed_mime,
        )
        _scan_ooxml_macros(data, ctx)
        try:
            from pptx import Presentation  # type: ignore
        except ImportError:
            bundle["errors"].append(
                issue("NeedsAdapter", "python-pptx not installed; install intake extras")
            )
            return _done(bundle, ctx, 0.0)

        try:
            prs = Presentation(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001
            bundle["errors"].append(issue("Corrupt", f"PPTX open failed: {exc}"))
            return _done(bundle, ctx, 0.0)

        if len(prs.slides) > ctx.limits.max_slides:
            bundle["errors"].append(
                issue(
                    "ResourceLimit",
                    f"PPTX has {len(prs.slides)} slides; max={ctx.limits.max_slides}",
                )
            )
            return _done(bundle, ctx, 0.0)

        slides_meta = []
        order = 0
        for si, slide in enumerate(prs.slides):
            texts: list[str] = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    texts.append(shape.text)
            body = "\n".join(texts)
            body, truncated = truncate_text(body, ctx.limits.max_segment_chars)
            if truncated:
                bundle["warnings"].append(issue("Truncated", f"slide {si + 1} truncated"))
            seg_id = f"slide-{si + 1}"
            bundle["segments"].append(
                {
                    "id": seg_id,
                    "kind": "text",
                    "text": body,
                    "order": order,
                    "locator": {"kind": "pptx_slide", "slide": si + 1},
                    "confidence": 0.9,
                }
            )
            slides_meta.append(
                {
                    "slide": si + 1,
                    "title": texts[0][:120] if texts else "",
                    "segment_ids": [seg_id],
                }
            )
            order += 1
        bundle["structure"]["slides"] = slides_meta
        bundle["confidence"] = 0.9 if bundle["segments"] else 0.3
        return _done(bundle, ctx, bundle["confidence"])


class XlsxAdapter:
    id = "xlsx"

    def supports(self, mime: str, filename: str | None, data: bytes) -> bool:
        if mime == XLSX_MIME:
            return True
        return (filename or "").lower().endswith(".xlsx")

    def extract(self, data: bytes, ctx: AdapterContext) -> dict[str, Any]:
        bundle = empty_bundle(
            file_id=ctx.file_id,
            content_hash=ctx.content_hash,
            mime_type=XLSX_MIME,
            size_bytes=ctx.size_bytes,
            path=ctx.path,
            path_mode=ctx.path_mode,
            adapter_id=self.id,
            sniffed_mime=ctx.sniffed_mime,
        )
        _scan_ooxml_macros(data, ctx)
        try:
            from openpyxl import load_workbook  # type: ignore
        except ImportError:
            bundle["errors"].append(
                issue("NeedsAdapter", "openpyxl not installed; install intake extras")
            )
            return _done(bundle, ctx, 0.0)

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            bundle["errors"].append(issue("Corrupt", f"XLSX open failed: {exc}"))
            return _done(bundle, ctx, 0.0)

        if len(wb.sheetnames) > ctx.limits.max_sheets:
            bundle["errors"].append(
                issue(
                    "ResourceLimit",
                    f"workbook has {len(wb.sheetnames)} sheets; max={ctx.limits.max_sheets}",
                )
            )
            wb.close()
            return _done(bundle, ctx, 0.0)

        sheets_meta = []
        order = 0
        for name in wb.sheetnames:
            ws = wb[name]
            seg_ids = []
            rows_preview: list[list[str]] = []
            for ri, row in enumerate(ws.iter_rows(values_only=True)):
                if ri >= ctx.limits.max_csv_rows:
                    bundle["warnings"].append(
                        issue("ResourceLimit", f"sheet {name} row limit reached")
                    )
                    break
                cells = [("" if c is None else str(c)) for c in row]
                rows_preview.append(cells)
                for ci, val in enumerate(cells):
                    if not val:
                        continue
                    from brainflow_worker.intake.adapters.structured import _col_letter

                    cell_ref = f"{_col_letter(ci)}{ri + 1}"
                    sid = f"{name}-{cell_ref}"
                    bundle["segments"].append(
                        {
                            "id": sid,
                            "kind": "cell",
                            "text": val,
                            "order": order,
                            "locator": {
                                "kind": "xlsx_cell",
                                "sheet": name,
                                "cell": cell_ref,
                            },
                            "confidence": 0.95,
                        }
                    )
                    seg_ids.append(sid)
                    order += 1
                    if order >= ctx.limits.max_segments:
                        bundle["warnings"].append(
                            issue("ResourceLimit", "XLSX segment limit")
                        )
                        break
                if order >= ctx.limits.max_segments:
                    break
            sheets_meta.append({"name": name, "segment_ids": seg_ids[:50]})
            if rows_preview:
                bundle["structure"].setdefault("tables", []).append(
                    {
                        "id": f"sheet-{name}",
                        "caption": name,
                        "rows": rows_preview[:100],
                        "locator": {"kind": "xlsx_cell", "sheet": name},
                    }
                )
            if order >= ctx.limits.max_segments:
                break
        wb.close()
        bundle["structure"]["sheets"] = sheets_meta
        bundle["confidence"] = 0.95 if bundle["segments"] else 0.3
        return _done(bundle, ctx, bundle["confidence"])


def _done(bundle: dict[str, Any], ctx: AdapterContext, confidence: float) -> dict[str, Any]:
    bundle["confidence"] = confidence
    apply_security(
        bundle,
        quarantine_actions=ctx.quarantine.actions,
        active_content_flags=ctx.quarantine.active_content_flags,
        macros_blocked=ctx.quarantine.macros_blocked,
        scripts_stripped=ctx.quarantine.scripts_stripped,
        limits_applied={
            "max_sheets": ctx.limits.max_sheets,
            "max_slides": ctx.limits.max_slides,
        },
    )
    return bundle
